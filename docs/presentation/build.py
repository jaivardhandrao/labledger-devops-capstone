"""Reproduce the reviewed 26 October deck from genuine repository evidence.

Published unchanged after student approval; its review-stage captions record
the original preparation. Current publication status is in demo/README.md.
"""
from pathlib import Path
from textwrap import dedent
import json
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
CLEANUP = json.loads((ROOT/'evidence/aws-live-2026-10-16/cloud-cleanup-verification.json').read_text())
assert CLEANUP['all_project_resources_removed'] is True, 'Verify actual teardown before building the final deck'
W, H = 960, 540
NAVY, INK, TEAL, PAPER, MUTED = '102D42', '17364B', '007E80', 'F4F8FA', '526C7B'
FONT = 'Helvetica'
BOLD = 'Helvetica-Bold'
font_dir = Path('/System/Library/Fonts/Supplemental')
if (font_dir/'Arial.ttf').exists():
    pdfmetrics.registerFont(TTFont('ArialLocal', str(font_dir/'Arial.ttf')))
    pdfmetrics.registerFont(TTFont('ArialLocal-Bold', str(font_dir/'Arial Bold.ttf')))
    FONT, BOLD = 'ArialLocal', 'ArialLocal-Bold'


def rgb(value):
    return tuple(int(value[i:i+2], 16)/255 for i in (0, 2, 4))


class Deck:
    def __init__(self):
        self.pdf = canvas.Canvas(str(OUT/'LabLedger_Capstone_Presentation.pdf'), pagesize=(W, H))
        self.pdf.setTitle('LabLedger — DevOps Capstone | Jaivardhan D Rao | Local and AWS Evidence')
        self.pdf.setAuthor('Jaivardhan D Rao')
        self.ppt = Presentation()
        self.ppt.slide_width, self.ppt.slide_height = Inches(13.333333), Inches(7.5)
        self.ppt.core_properties.title = 'LabLedger — DevOps Capstone'
        self.ppt.core_properties.author = 'Jaivardhan D Rao'
        self.number = 0
        self.items = []
        self.note_items = []

    def rect(self, x, y, w, h, color, radius=0):
        self.pdf.setFillColorRGB(*rgb(color))
        if radius:
            self.pdf.roundRect(x, H-y-h, w, h, radius, stroke=0, fill=1)
        else:
            self.pdf.rect(x, H-y-h, w, h, stroke=0, fill=1)
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                           Inches(x/72), Inches(y/72), Inches(w/72), Inches(h/72))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = 0.12

    def text(self, text, x, y, size=20, w=820, color=INK, bold=False, leading=1.25):
        font = BOLD if bold else FONT
        lines = []
        for block in text.split('\n'):
            line = ''
            for word in block.split():
                candidate = (line+' '+word).strip()
                if line and pdfmetrics.stringWidth(candidate, font, size) > w:
                    lines.append(line); line = word
                else:
                    line = candidate
            lines.append(line)
        h = size*leading*len(lines)
        assert x >= 0 and y >= 0 and x+w <= W+0.1 and y+h <= H+0.1, (self.number, text, y+h)
        self.pdf.setFillColorRGB(*rgb(color)); self.pdf.setFont(font, size)
        for i, line in enumerate(lines):
            self.pdf.drawString(x, H-y-size-i*size*leading, line)
        box = self.slide.shapes.add_textbox(Inches(x/72), Inches(y/72), Inches(w/72), Inches((h+3)/72))
        frame = box.text_frame; frame.clear(); frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line; p.font.name = 'Arial'; p.font.size = Pt(size); p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_before = p.space_after = Pt(0); p.line_spacing = Pt(size*leading)
        self.items.append({'slide': self.number, 'type': 'text', 'text': text, 'bounds': [x, y, w, h]})
        return h

    def image(self, relative, x, y, w, h):
        path = ROOT/relative
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w/iw, h/ih); dw, dh = iw*scale, ih*scale
        px, py = x+(w-dw)/2, y+(h-dh)/2
        self.pdf.drawImage(ImageReader(str(path)), px, H-py-dh, dw, dh, mask='auto')
        self.slide.shapes.add_picture(str(path), Inches(px/72), Inches(py/72), width=Inches(dw/72), height=Inches(dh/72))
        self.items.append({'slide': self.number, 'type': 'genuine_image', 'source': relative, 'bounds': [px, py, dw, dh]})

    def page(self, title, kicker, notes, dark=False):
        if self.number:
            self.pdf.showPage()
        self.number += 1
        self.slide = self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.rect(0, 0, W, H, NAVY if dark else PAPER)
        self.text(kicker.upper(), 64, 26, 11, color='A7E2DD' if dark else TEAL, bold=True)
        self.text(title, 64, 48, 30, color='FFFFFF' if dark else NAVY, bold=True)
        self.text(f'Jaivardhan D Rao  ·  10117  ·  Group A', 64, 512, 10,
                  color='CBDAE3' if dark else MUTED, w=600)
        self.text(f'{self.number:02d} / 12', 842, 512, 10, color='CBDAE3' if dark else MUTED, w=70)
        self.slide.notes_slide.notes_text_frame.text = dedent(notes).strip()
        self.note_items.append((title, dedent(notes).strip()))

    def bullet(self, heading, body, y, x=64, w=540):
        self.text(heading, x, y, 21, w=w, bold=True)
        return y+31+self.text(body, x, y+31, 17, w=w, color=MUTED)+18

    def link(self, label, url, x=64, y=478, w=820):
        self.text(label, x, y, 11, color=TEAL, w=w)
        self.pdf.linkURL(url, (x, H-y-16, x+w, H-y), relative=0, thickness=0)
        self.slide.shapes[-1].text_frame.paragraphs[0].runs[0].hyperlink.address = url

    def save(self):
        self.pdf.save()
        self.ppt.save(OUT/'LabLedger_Capstone_Presentation.pptx')
        (OUT/'layout-manifest.json').write_text(json.dumps({'slides': self.number, 'page_size_points': [W,H], 'items': self.items}, indent=2)+'\n')
        (OUT/'SPEAKER-NOTES.md').write_text('# LabLedger speaker notes\n\nLocal and actual temporary AWS evidence; captures and remaining limitations are identified.\n\n' + '\n\n'.join(f'## Slide {i}: {title}\n\n{notes}' for i,(title,notes) in enumerate(self.note_items,1)) + '\n')


d = Deck()
d.page('LabLedger', 'Review candidate · 26 October 2026', '''
Introduce the original campus equipment-loan domain. Identify yourself as Jaivardhan D Rao, roll 10117, Group A.
The screenshot is a genuine earlier School-profile capture of local Kubernetes, not AWS. Every displayed loan is fictional.
The application, CI, local Kubernetes and temporary personal AWS EKS ran successfully. Actual cleanup was verified separately before this final deck was built.
The title image remains the original local UI capture, not an AWS browser image. This 26 October review corrects documentation using saved 16 October evidence, with no fresh runtime checks. Private browser captures are not embedded. ArgoCD is not implemented and the complete slides-first presentation remains unrecorded.
''', dark=True)
d.text('Shared equipment.\nClear returns.', 64, 128, 32, w=318, color='FFFFFF', bold=True)
d.text('React · FastAPI · PostgreSQL\nDocker · CI · Helm · Monitoring', 64, 252, 19, w=305, color='CBDAE3')
d.rect(64, 355, 293, 86, '1E4559', radius=8)
d.text('VERIFIED ON EKS', 80, 370, 13, w=260, color='A7E2DD', bold=True)
d.text('Temporary cloud demo cleaned up.', 80, 398, 15, w=260, color='FFFFFF')
d.image('evidence/ui-kubernetes-desktop-8088.jpg', 390, 126, 515, 335)

d.page('One clear equipment-loan lifecycle', '01 · Application', '''
A borrowing request moves through requested, collected and returned. Only a collected loan past its due date is overdue.
A date has no UTC offset; the application evaluates today in Asia/Kolkata. Returning an item updates the statistics.
Explain the responsive table, keyboard access, filters and create/edit dialogs. The mobile capture is genuine and predates the later footer-only delivery change.
For the live demo open localhost:8088 only in the School Chrome profile; create and remove a temporary fictional loan only when that write is authorized.
''')
y = d.bullet('Request → collect → return', 'A lab desk records the borrower, equipment and return date.', 113, w=540)
y = d.bullet('Overdue has a precise meaning', 'Collected + due date before today in Asia/Kolkata. Requested loans are active, but not overdue.', y, w=540)
d.bullet('Usable on small screens', 'Search, status filters, readable actions and an accessible dialog. Demonstration records are explicitly fictional.', y, w=540)
d.image('evidence/ui-fixed-mobile-390.jpg', 664, 103, 220, 363)
d.text('Genuine School-profile mobile capture', 649, 472, 11, w=260, color=MUTED)

d.page('Follow a request through the stack', '02 · Architecture', '''
The browser sends relative /api requests. Compose Nginx resolves the backend service; Kubernetes Ingress routes /api directly to the API ClusterIP Service.
The frontend serves static React assets. The backend validates inputs and persists loans through SQLAlchemy in PostgreSQL, with Alembic owning schema migration.
Prometheus scrapes both backend replicas and Grafana queries Prometheus. Metrics Server is a separate HPA dependency.
The temporary cloud path was public HTTP NLB to NodePort30080, NGINX Ingress, ClusterIP Services and PostgreSQL. Monitoring remained loopback-only. The endpoint was retired after evidence capture.
''')
for i, (head, body) in enumerate([('React UI', 'Relative /api calls'), ('Nginx / Ingress', '/ and /api routes'), ('FastAPI', 'Validation + metrics'), ('PostgreSQL', 'Alembic + loans')]):
    x = 64+i*218
    d.rect(x, 132, 190, 90, 'FFFFFF', radius=8)
    d.text(head, x+14, 149, 19, w=164, bold=True)
    d.text(body, x+14, 183, 13, w=164, color=MUTED)
    if i<3: d.text('→', x+195, 158, 24, w=26, color=TEAL)
d.text('Compose', 64, 265, 20, bold=True, w=150)
d.text('localhost:3000 → frontend → backend → database', 234, 267, 18, w=650)
d.text('Kubernetes', 64, 313, 20, bold=True, w=150)
d.text('localhost:8088 → Ingress → ClusterIP Services → Pods', 234, 315, 18, w=650)
d.rect(64, 377, 832, 81, 'E4F1F2', radius=8)
d.text('Prometheus → Grafana', 82, 389, 20, w=400, color=TEAL, bold=True)
d.text('/metrics on both API replicas; real HTTP counters and latency histograms.', 82, 421, 17, w=780)

d.page('35 isolated tests, plus real-stack checks', '03 · Correctness', '''
The image is a native Terminal screenshot of the actual completed CI pytest output, part 2 of 2; it is explicitly not a fresh rerun. Both captures together show every test-result line.
The full evidence/pytest.txt contains all 35 passing API/migration tests. Tests isolate their SQLite database and mock dependencies; they do not use the running PostgreSQL data.
Mention invalid-input rejection, PUT replacement semantics, stats after return/delete, readiness failure and bounded metric route labels.
Separate real HTTP checks exercised the PostgreSQL/Compose path. The latest successful GitHub pipeline also ran the backend tests and nine frontend tests.
''')
d.image('evidence/presentation/terminal-ci-tests-part2.png', 64, 106, 832, 367)
d.text('Actual completed CI pytest output; all 35 results are available across the two native captures.', 64, 482, 12, color=MUTED)

d.page('Three healthy services; two non-root images', '04 · Docker and persistence', '''
This is an original actual local Terminal capture. The Compose service is named database, while the Kubernetes StatefulSet is named labledger-postgres. The later cloud database/table and container UID results are separate.
The API runs as UID10001 and frontend as UID101. The frontend is a Node-build to unprivileged-Nginx multi-stage image.
The read-only PostgreSQL session lists alembic_version and loans. No database record or schema was changed for this capture.
A prior forced backend IP-change check verified proxy DNS recovery and data preservation; see the original raw evidence. No images were rebuilt during this presentation capture.
''')
d.image('evidence/presentation/terminal-compose.png', 64, 106, 832, 367)
d.text('Live local containers and enforced read-only table inspection; original volumes preserved.', 64, 482, 12, color=MUTED)

d.page('A real change reached the running application', '05 · Delivery', '''
The displayed native Terminal image is historical cloud evidence: 06909ce, actual seven-green CI 38934060961, then exact-SHA EKS Helm release 2. It is not the final local revision.
The later local source change f271a837 has all seven jobs green in CI 38977458426. The genuine 3m54.5s recording shows a new accessibility refinement, 9 frontend tests/build, personal Git push, actual CI wait, anonymous pulls of the exact published ARM64 images, and local Helm release 4.
All four Ready local app Pods matched the published platform config digests. Separate School-profile captures show matching Arduino/requested filters and the new Show all loans reset.
The video is prepared privately under demo/ as a review candidate. It is Terminal-only and does not cover slides-first narration, AWS Console, ArgoCD or a full browser walkthrough. Delivery uses explicit Helm; no ArgoCD sync or automatic GitOps reconciliation is implemented.
''')
d.image('evidence/aws-live-2026-10-16/cloud-delivery-terminal.png', 64, 103, 832, 318)
d.text('Shown: historical 06909ce EKS rollout, release 2; infrastructure is retired.', 64, 431, 14, color=MUTED)
d.text('Later local: f271a837 → 7 green CI jobs → exact images → Helm release 4.', 64, 459, 14, color=TEAL, bold=True)
d.link('Open the final local change’s seven-job green run', 'https://github.com/jaivardhandrao/labledger-devops-capstone/actions/runs/38977458426', y=489)

d.page('The scan gate found a problem, then verified the repair', '06 · DevSecOps', '''
Explain that the initial Debian-based backend image failed strict HIGH/CRITICAL scanning. Its original failed report was retained.
Replacing the runtime with Alpine and rebuilding produced clean backend/frontend reports. Both amd64 and arm64 images must pass before canonical multiarch publication.
Trivy uses HIGH,CRITICAL with ignore-unfixed=false and exit-code1; no broad ignore gate was added. Earlier downloaded reports have zero HIGH/CRITICAL findings. Historical 06909ce and final local f271a837 metadata each verify all four strict scan jobs succeeded; their fresh raw report ZIPs were not downloaded.
A clean container scan is a bounded observation, not proof of total application security. Separate exact-credential and secret scans checked the public tree/history; credentials and local state stay outside Git.
''')
d.rect(64, 124, 393, 250, 'FFFFFF', radius=8)
d.text('Before', 86, 145, 16, color=MUTED, w=345)
d.text('Strict scan failed', 86, 181, 29, w=345, bold=True)
d.text('Debian runtime findings were preserved in the original report.', 86, 241, 20, w=338)
d.rect(481, 124, 415, 250, 'E4F1F2', radius=8)
d.text('After', 503, 145, 16, color=TEAL, w=365)
d.text('4 strict gates passed', 503, 181, 29, w=365, bold=True, color=TEAL)
d.text('Backend + frontend\namd64 + arm64\nNo HIGH / CRITICAL allowed', 503, 240, 20, w=365)
d.text('Gate: --severity HIGH,CRITICAL --ignore-unfixed=false --exit-code 1', 64, 408, 17, w=832, bold=True)
d.text('A clean scan is evidence about that image at that time; it is not a security guarantee.', 64, 456, 16, color=MUTED)

d.page('Helm release 2 ran on the dedicated EKS cluster', '07 · Kubernetes', '''
The screenshot displays the explicitly saved actual cloud API snapshot, with its original capture time and source file. It is not represented as a fresh cluster query after teardown. The existing local coursework cluster was preserved.
Both UI and API Deployments had two Ready replicas after HPA cooldown. Their Services are ClusterIP. PostgreSQL used an encrypted 2GiB gp3 PVC; the migration Job completed successfully.
A Completed migration Job is expected and does not mean application Pods are unhealthy. Probes gate application readiness on migration/database state.
Actual public NLB routing and cloud CRUD succeeded. This is genuine cloud evidence captured before approved teardown, not a claim that the temporary endpoint remains online.
''')
d.image('evidence/aws-live-2026-10-16/cloud-snapshot-terminal.png', 64, 106, 832, 367)
d.text('Actual EKS: ClusterIP Services, bound gp3, completed migration and deployed Helm releases.', 64, 482, 12, color=MUTED)

d.page('Resilience was exercised, not only configured', '08 · HPA and troubleshooting', '''
Use the actual cloud load/drill records. HPA rose from 2 to 4 Ready API Pods and returned to 2; the worker count remained one.
The 90-second burst recorded 7,431 HTTP200 and 12 failures. Failure categories were not retained, so their cause is not established. A separate lower-rate recovery check passed 40/40; this is not a claim that the burst workload was repaired.
A deliberately broken cloud Service selector produced502/no endpoints; exact restoration restored200. The initial EndpointSlice-null instrumentation error and its successful finally-restoration are retained separately.
Explain readiness vs liveness, Service selectors vs Pod labels, and why Metrics Server and CPU requests are required for HPA. HPA scales Pods, not worker nodes.
''')
d.text('2 → 4 Ready → 2', 64, 123, 44, color=TEAL, bold=True)
d.text('Observed cloud API HPA response; one worker throughout.', 64, 188, 20, color=MUTED)
d.rect(64, 245, 393, 139, 'FFFFFF', radius=8)
d.text('7,431 reads · 12 failures', 84, 264, 25, w=350, bold=True)
d.text('Separate lower-rate recovery: 40/40 HTTP 200. Burst failures retained.', 84, 311, 18, w=350, color=MUTED)
d.rect(481, 245, 415, 139, 'FFFFFF', radius=8)
d.text('200 → 502 → 200', 501, 264, 25, w=370, bold=True)
d.text('Service selector fault, diagnosis and real recovery.', 501, 311, 18, w=370, color=MUTED)
d.text('Actual cloud records retained. Earlier local 2,338 reads / 0 failures are separate evidence.', 64, 433, 16, color=MUTED)

d.page('Monitoring evidence has a clear scope', '09 · Observability', '''
This is a genuine native Terminal display of explicitly saved cloud checks, not a browser dashboard. Prometheus returned 16/16 UP before scaling; HPA temporarily increased the API target count.
Grafana's authenticated API returned the actual six-panel dashboard, six traffic series and positive samples. The genuine local six-panel browser capture was completed 16 October at 23:03:48 UTC after service recovery and remains private. Its one-hour range includes earlier revisions, rollout and recovery; it is not exclusively f271a837 or AWS evidence.
The initial local loopback-only control-plane failures and local repair remain historical evidence. EKS uses its separately documented managed-control-plane scrape profile.
Private Console/app/Prometheus browser images were captured separately and remain unpublished pending their specific publication approval. No browser pixels are embedded here.
''')
d.image('evidence/aws-live-2026-10-16/cloud-verification-terminal.png', 64, 104, 832, 335)
d.text('Cloud: 16/16 targets UP at capture time; actual CRUD and Service recovery retained.', 64, 449, 14, color=MUTED)
d.text('Local Grafana: six populated panels captured privately on 16 Oct, 23:03 UTC.', 64, 473, 14, color=TEAL, bold=True)
d.text('One-hour history spans revisions and recovery; private browser image is not embedded.', 64, 492, 11, color=MUTED)

d.page('Real AWS deployment, then verified cleanup', '10 · Terraform and lifecycle', '''
The actual frozen personal-account plan created 62 resources, zero changes/deletions. Mumbai EKS 1.34 and its managed t3.large group became ACTIVE. The VPC had two public/two private subnets, NAT and a Terraform-owned HTTP NLB.
The screenshot is an original native Terminal capture displaying actual completed destroy output and explicitly recorded post-deletion API checks. The separate apply capture preserves the actual 62-create output and ACTIVE EKS/node-group responses. Separate Console captures remain private until publication approval.
After actual runtime evidence, the disposable PostgreSQL PVC/EBS was removed while CSI/node permissions remained available. A separately reviewed 62-delete Terraform plan removed the new infrastructure; post-cleanup API/state checks are linked in the evidence gallery.
The final build requires a successful actual cleanup-verification file. The approved ceiling was US$10 / maximum two hours. The US$2.43 before-tax scenario is a planning estimate, not a measured final bill or automated billing cap.
''')
d.image('evidence/aws-live-2026-10-16/terraform-destroy-terminal.png', 64, 104, 832, 335)
d.text('62 created → real runtime verified → 62 destroyed; pre-existing resources preserved.', 64, 449, 15, color=TEAL, w=832, bold=True)
d.text('Planning estimate about US$2.43 before tax; actual final AWS bill is not yet verified.', 64, 477, 14, color=MUTED, w=832)

d.page('Present the verified work and its limits', '11 · Handoff', '''
This is a local review candidate, not a claim that the full assignment is ready. The 26 October 14:11 IST class announcement requires the video and PPT/PDF under demo/ in the same submitted repository by EOD; 14:17 IST also permits a Drive-video link in Markdown under demo/. No exact EOD clock time was specified in the supplied announcement.
M10 is README 2 marks, video 2, slides 1. The guide says a missing video yields zero for M10 only. Assessment remains the instructor's decision.
The 3m54.5s Terminal recording covers a genuine source change, push, CI and local deployment. It does not start with these slides or show the complete AWS Console, ArgoCD synced/GitOps, Kubernetes/Prometheus/Grafana browser flow. ArgoCD is not implemented; direct Helm deployment is not a substitute for it.
A genuine local populated Grafana screenshot exists privately, including earlier revisions and recovery. Saved AWS evidence documents a retired environment; no restart or provisioning is authorized. No new runtime checks were performed for this docs-only review.
The student reported manually submitting the form. No receipt, instructor acceptance, resubmission or new form action is claimed. Publishing the prepared demo/ and changed docs still requires specific approval.
''')
y = d.bullet('Prepared locally for review', 'Updated deck + actual local change-to-rollout video under demo/. Nothing new is published.', 116, w=820)
y = d.bullet('Full presentation remains incomplete', 'ArgoCD is absent. The slides-first AWS Console / GitOps / browser walkthrough is not recorded.', y, w=820)
d.bullet('Keep the evidence honest', 'Local populated Grafana capture is private historical evidence. AWS is retired; the form is user-reported submitted.', y, w=820)
d.link('Submitted repository; prepared review files are not yet published', 'https://github.com/jaivardhandrao/labledger-devops-capstone', y=480)
d.save()
print('Built 12-slide PDF and editable PPTX with speaker notes; genuine evidence images retained unchanged.')
